"""Dashboard, šabloni i podsetnici."""

from __future__ import annotations

from datetime import date, timedelta

import pytest

from app.db.models import TaskStatus
from app.services import budget as budget_service
from app.services import dashboard, reminders, templates
from app.services import dependencies as dep_service
from app.services import milestones as ms_service
from app.services import projects as projects_svc
from app.services import tasks as tasks_svc
from app.services import timetracking as tt
from scripts.seed_demo import seed

D = date(2026, 11, 2)


def test_dashboard_overview(session):
    p = seed(session, today=D)
    projects_svc.create(session, name="Arhiviran", archived=True)
    tt.add_entry(session, p.tasks[0].id, D, 2)
    ov = dashboard.build(session, today=D)
    assert [c.name for c in ov.projects] == [p.name]
    card = ov.projects[0]
    assert card.total == 8 and card.done == 2 and card.has_budget
    assert card.next_milestone == "Dizajn odobren"
    assert ov.open_tasks == 6 and ov.overdue_tasks == 0
    assert all(d.due <= D + timedelta(days=14) for d in ov.deadlines)
    assert ov.deadlines == sorted(ov.deadlines, key=lambda d: d.due)
    assert ov.hours_this_week >= 2
    assert dashboard.relative_day(0) == "danas" and dashboard.relative_day(-3) == "kasni 3 d"


def test_builtin_template_applies_with_dates_deps_and_budget(session):
    p = projects_svc.create(session, name="Novi sajt", start_date=D)
    n = templates.apply(session, p.id, templates.get_data(session, "builtin:sajt"), D)
    assert n == 10
    tree = tasks_svc.tree(session, p.id)
    first = tree[0][0]
    assert first.start_date == D and first.due_date == D + timedelta(days=3)
    assert any(depth == 1 for _, depth in tree)
    assert len(dep_service.for_project(session, p.id)) == 8
    assert len(ms_service.for_project(session, p.id)) == 2
    assert budget_service.summary(session, p.id).planned == 19700000


def test_save_project_as_template_and_reuse(session):
    src = seed(session, today=D)
    t = templates.save_from_project(session, src.id, "Moj šablon", "za picerije")
    infos = templates.list_templates(session)
    assert any(i.id == t.id and i.task_count == 10 and not i.builtin for i in infos)
    assert sum(1 for i in infos if i.builtin) == 3
    dst = projects_svc.create(session, name="Kopija")
    start = D + timedelta(days=30)
    templates.apply(session, dst.id, templates.get_data(session, t.id), start)
    src_titles = [x.title for x, _ in tasks_svc.tree(session, src.id)]
    dst_tree = tasks_svc.tree(session, dst.id)
    assert [x.title for x, _ in dst_tree] == src_titles
    assert all(x.status == TaskStatus.TODO for x, _ in dst_tree)  # stanje se ne kopira
    # datumi su pomereni za razliku početaka
    s0 = next(x for x, _ in tasks_svc.tree(session, src.id) if x.title == "Izrada sajta")
    d0 = next(x for x, _ in dst_tree if x.title == "Izrada sajta")
    assert (d0.start_date - s0.start_date).days == (start - src.start_date).days
    # isti naziv prepisuje, ugrađeni ne može da se obriše ni da se duplira
    assert templates.save_from_project(session, src.id, "Moj šablon").id == t.id
    with pytest.raises(ValueError):
        templates.save_from_project(session, src.id, "Izrada sajta")
    with pytest.raises(ValueError):
        templates.delete(session, "builtin:sajt")
    templates.delete(session, t.id)
    assert sum(1 for i in templates.list_templates(session) if not i.builtin) == 0


def test_reminders_once_per_day(session):
    p = projects_svc.create(session, name="P")
    tasks_svc.create(session, p.id, title="Danas", due_date=D)
    tasks_svc.create(session, p.id, title="Sutra", due_date=D + timedelta(days=1))
    tasks_svc.create(session, p.id, title="Kasni", due_date=D - timedelta(days=2))
    tasks_svc.create(session, p.id, title="Kasnije", due_date=D + timedelta(days=5))
    tasks_svc.create(session, p.id, title="Gotov", due_date=D, status=TaskStatus.DONE)
    ms_service.create(session, p.id, "Lansiranje", D + timedelta(days=1))
    items = reminders.pending(session, D)
    assert sorted((r.title, r.when) for r in items) == sorted(
        [
            ("Danas", "danas"),
            ("Sutra", "sutra"),
            ("Kasni", "kasni"),
            ("Milestone: Lansiranje", "sutra"),
        ]
    )
    title, body = reminders.summarize(items)
    assert title == "Rokovi: Kasni: 1, rok danas: 1, rok sutra: 2" and "Kasni" in body
    reminders.mark_sent(session, items, D)
    assert reminders.pending(session, D) == []
    assert len(reminders.pending(session, D + timedelta(days=1))) >= 2  # sledeći dan ponovo
    assert reminders.enabled(session)


def test_export_xlsx(session, tmp_path):
    from openpyxl import load_workbook

    from app.services import export_xlsx

    p = seed(session, today=D)
    tt.add_entry(session, p.tasks[0].id, D, 1.5, "sastanak")
    path = export_xlsx.write(session, p.id, tmp_path / "x.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Zadaci", "Troškovi i prihodi", "Budžet", "Vreme"]
    zad = wb["Zadaci"]
    assert zad.max_row == 11 and zad["D2"].value.date() == p.tasks[0].start_date
    budzet = wb["Budžet"]
    total = next(r for r in budzet.iter_rows(values_only=True) if r[0] == "Ukupno")
    assert total[1] == 222000 and abs(total[2] - 124930.8) < 0.001
    assert wb["Vreme"].max_row == 7
    assert any(r[2] == 1.5 for r in wb["Vreme"].iter_rows(min_row=2, values_only=True))
    costs = list(wb["Troškovi i prihodi"].iter_rows(min_row=2, values_only=True))
    assert any(r[1] == "Prihod" and r[8] == 120000 for r in costs)
