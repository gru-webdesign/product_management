"""Izvoz projekta u Excel (.xlsx): zadaci, troškovi i prihodi, budžet, vreme.

Datumi su pravi Excel datumi, iznosi pravi brojevi sa formatom, pa se mogu sortirati i sabirati.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.labels import PRIORITY, PROJECT_STATUS, TASK_STATUS
from app.services import budget as budget_service
from app.services import projects as project_service
from app.services import tasks as task_service
from app.services import timetracking as tt

HEAD_FILL = PatternFill("solid", fgColor="EAEFFD")
HEAD_FONT = Font(bold=True, color="18181C")
DATE_FMT = "DD.MM.YYYY"
MONEY_FMT = '#,##0.00 "{cur}"'
HOURS_FMT = "0.0#"


def _sheet(wb: Workbook, title: str, headers: list[str], widths: list[int], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 20
    return ws


def build(session: Session, project_id: str) -> Workbook:
    p = project_service.get(session, project_id)
    base = budget_service.base_currency(session)
    money = MONEY_FMT.format(cur=base)
    wb = Workbook()

    ws = _sheet(
        wb,
        "Zadaci",
        [
            "Naziv",
            "Status",
            "Prioritet",
            "Početak",
            "Rok",
            "Progres",
            "Procena (h)",
            "Utrošeno (h)",
            f"Planirani trošak ({base})",
            "Tagovi",
            "Opis",
        ],
        [44, 14, 12, 12, 12, 10, 12, 13, 20, 22, 50],
        first=True,
    )
    for t, depth in task_service.tree(session, project_id):
        ws.append(
            [
                ("    " * depth) + t.title,
                TASK_STATUS[t.status],
                PRIORITY[t.priority],
                t.start_date,
                t.due_date,
                t.progress / 100,
                t.estimate_hours,
                t.actual_hours,
                t.planned_cost_minor / 100 if t.planned_cost_minor else None,
                ", ".join(sorted(g.name for g in t.tags)),
                t.description,
            ]
        )
        r = ws.max_row
        ws.cell(r, 1).font = Font(bold=depth == 0)
        for c in (4, 5):
            ws.cell(r, c).number_format = DATE_FMT
        ws.cell(r, 6).number_format = "0%"
        for c in (7, 8):
            ws.cell(r, c).number_format = HOURS_FMT
        ws.cell(r, 9).number_format = money
    ws.auto_filter.ref = ws.dimensions

    ws = _sheet(
        wb,
        "Troškovi i prihodi",
        [
            "Datum",
            "Vrsta",
            "Kategorija",
            "Zadatak",
            "Opis",
            "Iznos",
            "Valuta",
            "Kurs",
            f"Iznos ({base})",
        ],
        [12, 10, 18, 30, 30, 14, 8, 10, 16],
    )
    for c in budget_service.costs(session, project_id):
        ws.append(
            [
                c.date,
                "Trošak",
                c.category.name,
                c.task.title if c.task else "",
                c.vendor or c.note,
                c.amount_minor / 100,
                c.currency,
                float(c.exchange_rate),
                -c.amount_base_minor / 100,
            ]
        )
    for i in budget_service.incomes(session, project_id):
        ws.append(
            [
                i.date,
                "Prihod",
                "",
                "",
                i.description,
                i.amount_minor / 100,
                i.currency,
                float(i.exchange_rate),
                i.amount_base_minor / 100,
            ]
        )
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 1).number_format = DATE_FMT
        ws.cell(r, 6).number_format = "#,##0.00"
        ws.cell(r, 8).number_format = "0.0000"
        ws.cell(r, 9).number_format = money
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions

    s = budget_service.summary(session, project_id)
    ws = _sheet(
        wb,
        "Budžet",
        ["Kategorija", "Planirano", "Potrošeno", "Preostalo", "Iskorišćeno"],
        [24, 16, 16, 16, 12],
    )
    for line in s.lines:
        ws.append(
            [
                line.name,
                line.planned / 100,
                line.spent / 100,
                line.remaining / 100,
                line.percent / 100,
            ]
        )
    ws.append(["Ukupno", s.planned / 100, s.spent / 100, s.remaining / 100, s.percent / 100])
    ws.append(["Prihod", s.income / 100])
    ws.append(["Zarada (prihod - potrošeno)", s.profit / 100])
    for r in range(2, ws.max_row + 1):
        for c in (2, 3, 4):
            ws.cell(r, c).number_format = money
        ws.cell(r, 5).number_format = "0%"
    total_row = len(s.lines) + 2
    for c in range(1, 6):
        ws.cell(total_row, c).font = Font(bold=True)

    ws = _sheet(
        wb,
        "Vreme",
        ["Datum", "Zadatak", "Sati", "Satnica", "Vrednost", "Napomena"],
        [12, 40, 10, 14, 16, 36],
    )
    for t, _ in task_service.tree(session, project_id):
        for e in tt.entries(session, t.id):
            ws.append(
                [
                    e.date,
                    t.title,
                    e.hours,
                    e.hourly_rate_minor / 100,
                    round(e.hours * e.hourly_rate_minor) / 100,
                    e.note,
                ]
            )
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 1).number_format = DATE_FMT
        ws.cell(r, 3).number_format = HOURS_FMT
        ws.cell(r, 4).number_format = money
        ws.cell(r, 5).number_format = money

    wb.properties.title = p.name
    wb.properties.subject = f"{p.client} {PROJECT_STATUS[p.status]}".strip()
    wb.properties.creator = "Etapa"
    return wb


def write(session: Session, project_id: str, path: Path) -> Path:
    build(session, project_id).save(path)
    return path
